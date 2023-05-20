import tkinter as tk
from tkinter import ttk
from MeasureSheet.ClassMeasureSheet import *
from MeasureFlameAngle.ClassMeasureFlameAngle import *
import openpyxl
from lib.lib_tkwidget import *
from lib.BaseClass import *
from lib.write_excel_mat import *
from lib.Table import *
class FrameChoosePic(BaseFrame):
    def __init__(self,master):
        self.father=master
        self.db=master.db
        self.back=self.db.Main.FrameChoosePic
        super().__init__(self.father,padding=5,relief='sunken')
        self.FramePanel=ttk.Frame(self)
        self.FramePanel.grid(row=0,column=0)
        self.create_panel()
        self.NotebookProcess=ttk.Notebook(self)
        self.NotebookProcess.grid(row=1,column=0)
        self.Table=TableFrame(self,[0],visible_rows=40)#TODO:调整默认宽度,切换液面就切换标题（绑定到每个子类）
        self.Table.grid(row=0,column=1,rowspan=2)
        #创建Notebook界面
        self.MeasureSheet=ClassMeasureSheet(self)
        self.NotebookProcess.add(self.MeasureSheet,text='MeasureSheet')
        self.MeasureFlameAngle=ClassMeasureFlameAngle(self)
        self.NotebookProcess.add(self.MeasureFlameAngle,text='MeasureFlameAngle')
        self.NotebookProcess.bind("<<NotebookTabChanged>>", self.on_page_changed)
        #Notebook初始化（选择第一个）
        self.NotebookProcess.select(0)
        self.on_page_changed()
    def on_page_changed(self,even=None):#切换页面
        self.Table.destroy()
        self.CurrentPage = getattr(self,self.NotebookProcess.tab(self.NotebookProcess.select(), "text"))
        self.Table=TableFrame(self,self.CurrentPage.back.TableHeadings,visible_rows=40)
        self.Table.grid(row=0,column=1,rowspan=2)
        print("Current page changed to:", self.CurrentPage)
    def create_panel(self):
        #R1
        self.create_var('Decimate',self.back,type='Int',init_value=0,FlagSave=True)
        self.LER1=LabelEntry(self.FramePanel,text='Decimate=',var=self.Decimate)
        self.LER1.grid(row=0,column=0,sticky='W')
        self.LER1.entry.bind('<Return>',self.CommandBindDecimate)
        ttk.Button(self.FramePanel,text='保存本组数据序列',command=self.CommandButtonSave).grid(row=0,column=1,sticky='W')
        #R2
        self.FrameR2C1=ttk.Frame(self.FramePanel)
        self.FrameR2C1.grid(row=1,column=0,sticky='W')
        self.create_var('FigNameStart',self.back,type='Int',init_value=0,FlagSave=True)
        self.create_var('FigNameEnd',self.back,type='Int',init_value=0,FlagSave=True)
        ttk.Label(self.FrameR2C1,text='图片范围：').pack(side='left')
        ttk.Label(self.FrameR2C1,textvariable=self.FigNameStart,width=5).pack(side='left')
        ttk.Label(self.FrameR2C1,text='-').pack(side='left')
        ttk.Label(self.FrameR2C1,textvariable=self.FigNameEnd,width=5).pack(side='left')
        self.FrameR2C2=ttk.Frame(self.FramePanel)
        self.FrameR2C2.grid(row=1,column=1,sticky='W')
        self.create_var('FigName',self.back,type='Int',init_value=0,FlagSave=True)
        ttk.Label(self.FrameR2C2,text='当前序号：').pack(side='left')
        Entry=ttk.Entry(self.FrameR2C2,textvariable=self.FigName,width=5)
        Entry.pack(side='left')
        Entry.bind('<Return>',self.CommandBindFigName)
        ttk.Button(self.FrameR2C2,text='+1',width=2,command=self.CommandButtonFigNameInc).pack(side='left')
        ttk.Button(self.FrameR2C2,text='-1',width=2,command=self.CommandButtonFigNameDec).pack(side='left')
        #R3
        self.create_var('ImpactX',self.back,type='Int',init_value=0,FlagSave=True)
        self.create_var('ImpactY',self.back,type='Int',init_value=0,FlagSave=True)
        self.FrameR3C1=ttk.Frame(self.FramePanel)
        self.FrameR3C1.grid(row=2,column=0,sticky='W')
        ttk.Label(self.FrameR3C1,text='基点：').pack(side='left')
        Entry=ttk.Entry(self.FrameR3C1,textvariable=self.ImpactX,width=4)
        Entry.pack(side='left')
        Entry.bind('<Return>',self.CommandBindImpactXY)
        ttk.Label(self.FrameR3C1,text=',').pack(side='left')
        Entry=ttk.Entry(self.FrameR3C1,textvariable=self.ImpactY,width=4)
        Entry.pack(side='left')
        Entry.bind('<Return>',self.CommandBindImpactXY)
        self.create_var('PixSize',self.back,type='Double',init_value=0,FlagSave=True)
        LE=LabelEntry(self.FramePanel,'像素尺寸=',self.PixSize,width=6,extratext='um/pix')
        LE.grid(row=2,column=1,sticky='W')
        LE.entry.bind('<Return>',self.CommandBindPixSize)
    def CommandBindDecimate(self,*args):
        self.back.update('Decimate',self.Decimate.get(),False)
    
    def FigNameFresh1(self):
        self.back.save_data(self.CurrentPage.back.generate_data())
        self.Table.update_table(self.back.data[:])
    def CommandButtonFigNameInc(self):
        self.FigNameFresh1()
        if self.FigName.get()+1 in self.back.get('FigureNameList'):
            self.back.update('FigName',self.FigName.get()+1)
            self.Fresh()
    def CommandButtonFigNameDec(self):
        self.FigNameFresh1()
        if self.FigName.get()-1 in self.back.get('FigureNameList'):
            self.back.update('FigName',self.FigName.get()-1)
            self.Fresh()
    def CommandBindFigName(self,*args):
        self.FigNameFresh1()
        if self.FigName.get() in self.back.get('FigureNameList'):
            self.back.update('FigName',self.FigName.get())
            self.Fresh()
    def CommandBindImpactXY(self,*args):
        self.back.update('ImpactX',self.ImpactX.get(),False)
        self.back.update('ImpactY',self.ImpactY.get(),False)
        #后端修改
        self.Fresh()
    def CommandBindPixSize(self,*args):
        self.back.update('PixSize',self.PixSize.get(),False)
        self.Fresh()
    def CommandButtonSave(self):
        #写入自己的表
        if not self.father.back.FlagDirect:
            wb_target=openpyxl.Workbook()
            ws_target=wb_target.active
            title=[self.CurrentPage.back.generate_title(),
                self.CurrentPage.back.TableHeadings]
            write_excel_mat(ws_target,title)
            write_excel_mat(ws_target,self.back.data,FlagContinue=True)
            wb_target.save(self.back.get_dir_res(self.CurrentPage.__class__.__name__))
        #写入汇总表
        wb_target=openpyxl.load_workbook(self.father.back.FDirRes)
        ws_target=wb_target.active
        for i in range(1,ws_target.max_row+1):
            if ws_target.cell(i,1).value==self.father.back.generate_ImageLabel():
                cache=self.CurrentPage.back.generate_title()
                for k in range(len(cache)//2):
                    for j in range(1,ws_target.max_column+1):
                        if ws_target.cell(2,j).value==cache[k*2]:
                            ws_target.cell(i,j,cache[k*2+1])
        wb_target.save(self.father.back.FDirRes)
    def Fresh(self):
        self.CurrentPage.FreshSelf()
if __name__=='__main__':
    root=tk.Tk()
    App=FrameChoosePic(root)
    App.pack()
    # root.protocol("WM_DELETE_WINDOW",close_window)
    root.mainloop()