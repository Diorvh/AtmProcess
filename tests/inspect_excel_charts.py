"""Run explicitly to inspect exported charts in a separate, hidden Excel instance."""

import argparse
import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from AtmProcess.core.excel_workbook_coordinator import ExcelWorkbookCoordinator
from AtmProcess.domain.settings import AtomizationSettings
from AtmProcess.services.atomization_angle_service import AtomizationBatchResult
from AtmProcess.services.atomization_result_export_service import AtomizationExportConfig, AtomizationResultExportService


def inspect_charts(frame_count: int, max_freq: int, verify: bool = False) -> None:
    import comtypes.client
    from comtypes import COMError

    coordinator = ExcelWorkbookCoordinator(active_excel_provider=lambda: None, file_opener=lambda _path: None)
    exporter = AtomizationResultExportService(coordinator)
    frames = np.arange(frame_count)
    left = 30 + 0.01 * frames
    right = 150 - 0.02 * frames
    result = AtomizationBatchResult(np.column_stack((frames, left, right, right - left)), 0.1, 1)
    with TemporaryDirectory(prefix="atm-chart-check-") as directory:
        root = Path(directory)
        config = AtomizationExportConfig(root, root / "absent.xlsx", max_freq, "test", output_dir=root)
        path = exporter.export(config, AtomizationSettings(), result)
        saved = openpyxl.load_workbook(path)
        try:
            chart_ranges = [
                [{"x": series.cat.numRef.f, "y": series.val.numRef.f} for series in chart.series]
                for chart in saved.active._charts
            ]
            frequency_rows = sum(saved.active.cell(row, 6).value is not None for row in range(5, saved.active.max_row + 1))
        finally:
            saved.close()

        excel = comtypes.client.CreateObject("Excel.Application", dynamic=True)  # Never attach to the user's Excel session.
        workbook = None
        try:
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = excel.Workbooks.Open(str(path), 0, True)
            excel.CalculateFull()
            charts = workbook.Worksheets(1).ChartObjects()
            report = []
            for index in range(1, charts.Count + 1):
                chart = charts.Item(index).Chart
                chart.Refresh()
                series_data = []
                for series_index in range(1, chart.SeriesCollection().Count + 1):
                    series = chart.SeriesCollection().Item(series_index)
                    data = {
                        "name": series.Name,
                        "formula": series.Formula,
                    }
                    for label, property_name in (("x", "XValues"), ("y", "Values")):
                        try:
                            values = getattr(series, property_name)
                            values = values if isinstance(values, tuple) else (values,)
                            data[label] = {"count": len(values), "first": values[0], "last": values[-1]}
                        except COMError as error:
                            data[label] = {"error": str(error)}
                    series_data.append(data)
                report.append({"chart": index, "saved_ranges": chart_ranges[index - 1], "excel_series": series_data})
            print(json.dumps({"frames": frame_count, "frequency_rows": frequency_rows, "charts": report}, ensure_ascii=False, indent=2))
            if verify:
                assert len(report) == (3 if frequency_rows else 1)
                for chart_index, item in enumerate(report):
                    expected_count = frame_count if chart_index == 0 else frequency_rows
                    expected_last_x = frame_count - 1 if chart_index == 0 else (frequency_rows - 1) * 40000 / frame_count
                    assert len(item["excel_series"]) == 4
                    for series_index, data in enumerate(item["excel_series"]):
                        assert data["name"] == ("左锥角", "右锥角", "雾化锥角", "锥角中心线")[series_index]
                        assert data["x"]["count"] == data["y"]["count"] == expected_count
                        assert data["x"]["first"] == 0
                        assert abs(data["x"]["last"] - expected_last_x) < 1e-8
                        expected_first_y = (30, 150, 120, 90)[series_index] if chart_index == 0 else 0
                        assert abs(data["y"]["first"] - expected_first_y) < 1e-8
                        assert data["y"]["last"] is not None
                print("Excel native chart checks passed.")
        finally:
            try:
                if workbook is not None:
                    workbook.Close(False)
            finally:
                excel.Quit()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--frames", type=int, default=1000)
    parser.add_argument("--max-freq", type=int, default=100)
    parser.add_argument("--verify", action="store_true")
    arguments = parser.parse_args()
    inspect_charts(arguments.frames, arguments.max_freq, arguments.verify)
