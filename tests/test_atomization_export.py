import shutil
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from AtmProcess.application.actions import ProcessingActions
from AtmProcess.core.excel_workbook_coordinator import ExcelWorkbookCoordinator
from AtmProcess.domain.context import ExperimentContext, FrameSelection
from AtmProcess.domain.settings import AppSettings, AtomizationSettings
from AtmProcess.services.atomization_angle_service import AtomizationBatchResult
from AtmProcess.services.atomization_result_export_service import AtomizationExportConfig, AtomizationResultExportService


class AtomizationExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.coordinator = ExcelWorkbookCoordinator(
            active_excel_provider=lambda: None,
            file_opener=lambda _path: None,
        )
        self.exporter = AtomizationResultExportService(self.coordinator)
        self.settings = AtomizationSettings(y=107, r_start_left=12, r_start_right=34)
        self.result = AtomizationBatchResult(
            rows=np.array([[0, 30, 150, 120], [4, 32, 148, 116]], dtype=float),
            compute_time=0.1,
            process_count=1,
        )

    def source_sheet(self):
        workbook = openpyxl.Workbook()
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.exporter._write_batch_data(sheet, self.settings, self.result)
        self.exporter._write_angle_mid(sheet, np.mean(self.result.rows[:, 1:3], axis=1))
        return sheet

    def summary_workbook(self, source, paired_columns=False):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        headers = [source.cell(1, column).value for column in range(1, source.max_column + 1)]
        if not paired_columns:
            headers = [title for title in headers if title not in ("RStartL", "RStartR")]
        sheet.append(["编号", "Y", "备注", *headers, "AngleL", "AngleLStd"])
        sheet.append(["1pix", 800, "calibration"])
        sheet.append([7, 700, "other experiment"])
        sheet.append([24, 600, "target experiment"])
        path = self.root / "summary.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def test_settings_preserve_frame_rate_and_legacy_default(self):
        settings = AppSettings()
        settings.atomization.frame_rate_hz = 12500
        self.assertEqual(AppSettings.from_dict(settings.to_dict()).atomization.frame_rate_hz, 12500)
        self.assertEqual(AppSettings.from_dict({}).atomization.frame_rate_hz, 40000)

    def test_frequency_axis_uses_supplied_rate_directly(self):
        for rate in (10000, 20000):
            with self.subTest(rate=rate):
                workbook = openpyxl.Workbook()
                self.addCleanup(workbook.close)
                signal = np.sin(2 * np.pi * np.arange(100) / 10)
                matrix = np.column_stack([signal] * 4)
                self.exporter._write_frequency_data(workbook.active, matrix, 50, rate)
                for column in range(7, 11):
                    peak_row = max(range(6, 55), key=lambda row: workbook.active.cell(row, column).value)
                    self.assertEqual(workbook.active.cell(peak_row, 6).value, rate / 10)

    def test_result_adds_only_radius_start_columns(self):
        sheet = self.source_sheet()
        self.assertEqual(sheet.max_column, 16)
        self.assertEqual([sheet.cell(1, col).value for col in range(9, 13)], ["NumAvgL", "RStartL", "NumAvgR", "RStartR"])
        self.assertEqual([sheet.cell(2, col).value for col in (10, 12)], [12, 34])

    def test_summary_supports_radius_columns_next_to_each_num_avg(self):
        source = self.source_sheet()
        path = self.summary_workbook(source, paired_columns=True)
        config = AtomizationExportConfig(self.root, path, 100, "实验编号24", experiment_no=24)
        self.exporter._update_summary(config, source)
        workbook = openpyxl.load_workbook(path)
        try:
            sheet = workbook.active
            self.assertEqual(sheet.max_column, 21)
            self.assertEqual(sheet.cell(4, 2).value, 600)
            for column in range(1, 17):
                self.assertAlmostEqual(sheet.cell(4, column + 3).value, source.cell(2, column).value)
        finally:
            workbook.close()
        actions = ProcessingActions(SimpleNamespace(experiment=ExperimentContext(experiment_no=24, result_summary=path)))
        parameters = actions.read_atomization_parameters_from_summary()
        self.assertEqual(parameters["r_start_left"], 12)
        self.assertEqual(parameters["r_start_right"], 34)
        self.assertEqual(parameters["num_avg_right"], self.settings.num_avg_right)
        self.assertEqual(parameters["thres_ratio_left"], self.settings.thres_ratio_left)

    def test_sample_summary_accepts_batch_results_and_parameter_readback(self):
        template = Path(__file__).resolve().parents[1] / "sample_data/atomization_summary_template.xlsx"
        path = self.root / "summary.xlsx"
        shutil.copy2(template, path)
        config = AtomizationExportConfig(self.root, path, 100, "实验编号29", experiment_no=29)
        self.exporter.export(config, self.settings, self.result)
        workbook = openpyxl.load_workbook(path)
        try:
            sheet = workbook.active
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(workbook.sheetnames, ["测量数据"])
            self.assertEqual(sheet.max_row, 30)
            self.assertEqual(sheet.cell(30, headers["编号"]).value, 29)
            self.assertTrue(all(sheet.cell(row, headers["编号"]).value is None for row in range(2, 30)))
            self.assertIsNone(sheet["B2"].value)
            self.assertNotIn("图像备注", headers)
            self.assertEqual(sheet.cell(30, headers["U"]).value, 26.19499)
            self.assertEqual(sheet.cell(30, headers["Pc"]).value, 4.99628)
            self.assertEqual(sheet["F30"].value, "=E30*40")
            self.assertEqual(sheet["Y30"].value, "=X30+$X$2")
            self.assertNotIn("平均破碎长度/mm", headers)
            for index, title in enumerate(("AngleL", "AngleR", "AtmAngle", "AngleM")):
                values = self.result.rows[:, index + 1] if index < 3 else np.mean(self.result.rows[:, 1:3], axis=1)
                self.assertAlmostEqual(sheet.cell(30, headers[title]).value, np.mean(values))
                self.assertAlmostEqual(sheet.cell(30, headers[title] + 1).value, np.std(values, ddof=1))
        finally:
            workbook.close()
        actions = ProcessingActions(SimpleNamespace(experiment=ExperimentContext(experiment_no=29, result_summary=path)))
        parameters = actions.read_atomization_parameters_from_summary()
        self.assertEqual(parameters["y"], self.settings.y)
        self.assertEqual(parameters["r_start_left"], self.settings.r_start_left)
        self.assertEqual(parameters["r_start_right"], self.settings.r_start_right)

    def test_summary_uses_experiment_id_and_keeps_other_parameter_blocks(self):
        source = self.source_sheet()
        path = self.summary_workbook(source)
        config = AtomizationExportConfig(self.root, path, 100, "custom label", experiment_no=24)
        for _ in range(2):
            self.exporter._update_summary(config, source)
        workbook = openpyxl.load_workbook(path)
        try:
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet.max_column, 21)
            self.assertEqual(sheet.cell(4, 2).value, 600)
            self.assertEqual(sheet.cell(4, 6).value, self.settings.y)
            self.assertIsNone(sheet.cell(3, 6).value)
            np.testing.assert_allclose(
                [sheet.cell(4, col).value for col in (18, 19)],
                [source.cell(4, 2).value, source.cell(5, 2).value],
            )
            self.assertEqual([sheet.cell(1, col).value for col in (20, 21)], ["RStartL", "RStartR"])
            self.assertEqual([sheet.cell(4, col).value for col in (20, 21)], [12, 34])
        finally:
            workbook.close()

        actions = ProcessingActions(SimpleNamespace(experiment=ExperimentContext(experiment_no=24, result_summary=path)))
        parameters = actions.read_atomization_parameters_from_summary()
        self.assertEqual(parameters["r_start_left"], 12)
        self.assertEqual(parameters["r_start_right"], 34)
        self.assertEqual(parameters["y"], self.settings.y)

    def test_legacy_summary_does_not_replace_missing_radius_start_values(self):
        path = self.summary_workbook(self.source_sheet())
        actions = ProcessingActions(SimpleNamespace(experiment=ExperimentContext(experiment_no=24, result_summary=path)))
        self.exporter.summary_service.update_contiguous_block(path, "实验编号24", "XLeft", 0, [123], 24)
        parameters = actions.read_atomization_parameters_from_summary()
        self.assertEqual(parameters["x_left"], 123)
        self.assertNotIn("r_start_left", parameters)
        self.assertNotIn("r_start_right", parameters)

    def test_batch_passes_frame_rate_without_decimation_adjustment(self):
        settings = AppSettings()
        settings.experiment = ExperimentContext(
            data_root=self.root,
            experiment_no=24,
            frames=FrameSelection(start=0, end=8, decimate=4),
        )
        settings.atomization.frame_rate_hz = 12345
        actions = ProcessingActions(SimpleNamespace(settings=settings, experiment=settings.experiment))
        actions.atomization_service = Mock()
        actions.atomization_export_service = Mock()
        actions.run_atomization_batch()
        self.assertEqual(actions.atomization_service.run_batch.call_args.args[1], [0, 4, 8])
        config, atomization, _result = actions.atomization_export_service.export.call_args.args
        self.assertEqual(config.experiment_no, 24)
        self.assertEqual(atomization.frame_rate_hz, 12345)


if __name__ == "__main__":
    unittest.main()
