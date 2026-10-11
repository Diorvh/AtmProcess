import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from AtmProcess.core.excel_workbook_coordinator import ExcelWorkbookCoordinator
from AtmProcess.services.measurement_export_service import MeasurementExportService
from AtmProcess.ui.pages.processing_pages import MeasurementPage


class MeasurementStatisticsTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        coordinator = ExcelWorkbookCoordinator(active_excel_provider=lambda: None, file_opener=lambda _path: None)
        self.exporter = MeasurementExportService(workbook_coordinator=coordinator)

    def test_breakup_gui_export_summary_and_import_use_recorded_point_mean(self):
        for frames, lengths in (([0, 2, 4], [1, 10, 1]), ([0, 1, 100], [1, 10, 1]), ([8], [7])):
            with self.subTest(frames=frames):
                summary = self.root / "summary.xlsx"
                workbook = openpyxl.Workbook()
                workbook.active.append(["编号", "X", "Y", "像素大小/um/pix", "平均破碎长度/mm", "Decimate", "总张数", "起始帧", "结束帧"])
                workbook.active.append([99, 555])
                workbook.active.append([24])
                workbook.save(summary)
                workbook.close()
                rows = [[str(frame), frame, length, length, length] for frame, length in zip(frames, lengths)]
                mean = float(np.mean(lengths))
                std = float(np.std(lengths, ddof=1)) if len(lengths) > 1 else 0.0
                path = self.exporter.export_breakup(self.root, "实验编号24", rows, 0, 0, 1000, 2,
                                                    output_dir=self.root, summary_path=summary, experiment_no=24)
                exported = openpyxl.load_workbook(path, data_only=True)
                try:
                    self.assertEqual(exported.active.cell(3, 6).value, mean)
                    self.assertAlmostEqual(exported.active.cell(3, 7).value, std)
                finally:
                    exported.close()
                recorded = openpyxl.load_workbook(summary, data_only=True)
                try:
                    self.assertEqual(recorded.active.cell(3, 5).value, mean)
                    self.assertEqual(recorded.active.cell(2, 2).value, 555)
                finally:
                    recorded.close()
                imported = self.exporter.read_records(path)
                for displayed_rows in (rows, imported.rows):
                    view = SimpleNamespace(mode=SimpleNamespace(get=lambda: "breakup"),
                                           breakup_records=displayed_rows, record_stats_text=Mock())
                    MeasurementPage._refresh_record_stats(view)
                    displayed = view.record_stats_text.set.call_args.args[0]
                    self.assertIn(f"平均值：{mean:.4g} mm", displayed)
                    self.assertIn(f"标准差：{std:.4g} mm", displayed)

    def test_sheet_length_and_width_keep_arithmetic_mean(self):
        widths, lengths = [1, 10, 1], [2, 20, 2]
        rows = [[str(frame), frame, 0, width, 0, length, width, width, length, length]
                for frame, width, length in zip([0, 2, 4], widths, lengths)]
        path = self.exporter.export_sheet(self.root, "实验编号24", rows, 0, 0, 1000, output_dir=self.root)
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            self.assertEqual(workbook.active.cell(3, 11).value, 4)
            self.assertEqual(workbook.active.cell(3, 13).value, 8)
        finally:
            workbook.close()
        view = SimpleNamespace(mode=SimpleNamespace(get=lambda: "sheet"), sheet_records=rows, record_stats_text=Mock())
        MeasurementPage._refresh_record_stats(view)
        displayed = view.record_stats_text.set.call_args.args[0]
        self.assertIn("宽度平均值：4 mm", displayed)
        self.assertIn("长度平均值：8 mm", displayed)


if __name__ == "__main__":
    unittest.main()
