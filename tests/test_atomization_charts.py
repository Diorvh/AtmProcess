import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import numpy as np
import openpyxl
from openpyxl.utils.cell import range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from AtmProcess.core.excel_workbook_coordinator import ExcelWorkbookCoordinator
from AtmProcess.domain.settings import AtomizationSettings
from AtmProcess.services.atomization_angle_service import AtomizationBatchResult
from AtmProcess.services.atomization_result_export_service import AtomizationExportConfig, AtomizationResultExportService


class AtomizationChartTests(unittest.TestCase):
    def test_saved_charts_keep_all_points_and_labels(self):
        coordinator = ExcelWorkbookCoordinator(active_excel_provider=lambda: None, file_opener=lambda _path: None)
        exporter = AtomizationResultExportService(coordinator)
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for count, limit in ((1000, 100), (10000, 100), (9, 3), (10, 100), (2, 1), (1, 100), (10, 0)):
                with self.subTest(frames=count, max_freq=limit):
                    frames = np.arange(count)
                    left = 30 + 0.01 * frames
                    right = 150 - 0.005 * frames
                    result = AtomizationBatchResult(np.column_stack((frames, left, right, right - left)), 0.1, 1)
                    config = AtomizationExportConfig(root, root / "absent.xlsx", limit, "test", output_dir=root)
                    path = exporter.export(config, AtomizationSettings(), result)
                    workbook = openpyxl.load_workbook(path)
                    try:
                        sheet = workbook.active
                        frequency_count = min(count // 2, limit)
                        self.assertEqual(len(sheet._charts), 3 if frequency_count else 1)
                        for index, chart in enumerate(sheet._charts):
                            first_row = 7 if index == 0 else 5
                            point_count = count if index == 0 else frequency_count
                            x_column = 1 if index == 0 else 6
                            self.assertEqual(len(chart.series), 4)
                            for series_index, series in enumerate(chart.series):
                                last_row = first_row + point_count - 1
                                x_bounds = range_boundaries(series.cat.numRef.f.split("!")[1])
                                y_bounds = range_boundaries(series.val.numRef.f.split("!")[1])
                                self.assertEqual(x_bounds, (x_column, first_row, x_column, last_row))
                                column = x_column + series_index + 1
                                self.assertEqual(y_bounds, (column, first_row, column, last_row))
                                self.assertEqual(series.tx.v, ("左锥角", "右锥角", "雾化锥角", "锥角中心线")[series_index])
                                self.assertIsNotNone(sheet.cell(first_row, column).value)
                                self.assertIsNotNone(sheet.cell(last_row, column).value)
                        self.assertEqual(sheet.cell(7, 2).value, left[0])
                        self.assertAlmostEqual(sheet.cell(count + 6, 2).value, left[-1])
                    finally:
                        workbook.close()


if __name__ == "__main__":
    unittest.main()
