import unittest
from pathlib import Path

import openpyxl

from algorithms.breakup import calculate_breakup_length
from algorithms.sheet import calculate_sheet_measurement
from core.atomization_angle_runner import AtomizationAngleBatchConfig
from core.excel_io import write_matrix
from core.image_io import read_image
from core.parallel import split_frame_range
from core.paths import experiment_dir, image_path


class CoreTests(unittest.TestCase):
    def test_atomization_legacy_input_format(self):
        config = AtomizationAngleBatchConfig(
            data_dir="data",
            fig_start=10,
            fig_end=20,
            decimate=2,
            core_num=4,
            x_left=1,
            x_right=2,
            y=3,
            r_max=4,
            theta_left=5,
            theta_right=6,
            num_avg_left=7,
            r_start_left=8,
            num_avg_right=9,
            r_start_right=10,
            thres_ratio_left=0.5,
            thres_ratio_right=0.6,
            window_length=101,
            polyorder=1,
            save_marked_images=True,
        )
        lines = config.to_legacy_lines()
        self.assertEqual(len(lines), 21)
        self.assertEqual(lines[0], "data")
        self.assertEqual(lines[-2:], ["1", "1"])

    def test_sheet_measurement(self):
        result = calculate_sheet_measurement(10, 30, 5, 45, 20)
        self.assertEqual(result.width_px, 20)
        self.assertAlmostEqual(result.width_mm, 0.4)
        self.assertEqual(result.length_px, 40)
        self.assertAlmostEqual(result.length_mm, 0.8)

    def test_breakup_length(self):
        result = calculate_breakup_length(75, 58, 20)
        self.assertEqual(result.length_px, 17)
        self.assertAlmostEqual(result.length_mm, 0.34)

    def test_paths(self):
        self.assertEqual(experiment_dir("root", 3), Path("root") / "实验编号3")
        self.assertEqual(image_path("root", 100), Path("root") / "100.tif")

    def test_excel_write_matrix(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        write_matrix(ws, [[1, 2], [3, 4]])
        self.assertEqual(ws.cell(1, 1).value, 1)
        self.assertEqual(ws.cell(2, 2).value, 4)

    def test_read_image_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            read_image("__missing__.tif")

    def test_split_frame_range_clamps_part_count(self):
        chunks = split_frame_range(10, 12, 8)
        self.assertEqual(chunks, [[10], [11], [12]])

    def test_split_frame_range_balances_work(self):
        chunks = split_frame_range(1, 5, 2)
        self.assertEqual(chunks, [[1, 2, 3], [4, 5]])


if __name__ == "__main__":
    unittest.main()
