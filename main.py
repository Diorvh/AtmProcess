import tkinter as tk
from tkinter import ttk
from database import db
from lib_tkwidget import *
from ClassProcess import ClassProcess
import matplotlib.pyplot as plt
import os
import numpy as np
import openpyxl
#依赖库：pip install numpy TinyDB matplotlib openpyxl
class QProcess(ttk.Frame):
    def __init__(self,master):
        super().__init__(master)
        self.digit=5
        self.create_panel()
        self.create_res()
        self.Process=[]
        self.Process.append(ClassProcess(self,name='整体',interval=True,digit=self.digit))
        self.Process[-1].grid(row=1,column=0)
        self.Process.append(ClassProcess(self,name='实验段',interval=[0,2000],digit=self.digit))
        self.Process[-1].grid(row=1,column=1)
        self.Process.append(ClassProcess(self,name='计算cd',interval=[6000,7000],digit=self.digit))
        self.Process[-1].grid(row=1,column=2)
        self.fresh_data()
    def create_panel(self):
        self.FramePanel=ttk.Frame(self,padding=5,relief='sunken')
        self.FramePanel.grid(row=0,column=0,columnspan=2,sticky='NSWE')
        self.new_var('NumPipe','Int')
        self.FrameR0=ttk.Frame(self.FramePanel)
        self.FrameR0.pack(fill='x')
        ttk.Label(self.FrameR0,text='管路数量：').pack(side='left')
        self.combobox=ttk.Combobox(self.FrameR0,textvariable=self.NumPipe,values=[1,2])
        self.combobox.pack(side='left')
        self.new_var('DirData')
        ChooseFile(self.FramePanel,'数据文件夹：',self.DirData,type='Dir',width=100).pack()
        self.FrameR1=ttk.Frame(self.FramePanel)
        self.FrameR1.pack(fill='x')
        self.new_var('ExpNo','Int')
        LabelEntry(self.FrameR1,'实验编号',self.ExpNo).pack(side='left')
        ttk.Button(self.FrameR1,text='+1',width=2,command=self.CommandButtonExpNoInc).pack(side='left')
        ttk.Button(self.FrameR1,text='-1',width=2,command=self.CommandButtonExpNoDec).pack(side='left')
        ttk.Button(self.FrameR1,text='刷新',width=4,command=self.CommandButtonReadFDir).pack(side='left')
        self.new_var('FDirData')
        self.new_var('FDirRes')
        ChooseFile(self.FramePanel,'数据文件：',self.FDirData,width=100).pack()
        ChooseFile(self.FramePanel,'结果文件：',self.FDirRes,width=100).pack()
    def create_res(self):
        self.FrameRes=ttk.Frame(self,padding=5,relief='sunken')
        self.FrameRes.grid(row=0,column=2,sticky='WE')
        self.FrameR1=ttk.Frame(self.FrameRes)
        self.FrameR1.pack()
        self.VarName=['rou','D','cd','Pccd','DPcd','U','Pc']
        self.VarName_unit=['rou','mm','','MPa','MPa','m/s','MPa']
        for i in self.VarName:
            setattr(self,'Data'+i,tk.DoubleVar(self))
            setattr(self,'DataD'+i,tk.DoubleVar(self))
        ttk.Label(self.FrameR1,text='实验编号').pack(side='left')
        ttk.Label(self.FrameR1,textvariable=self.ExpNo).pack(side='left')
        ttk.Label(self.FrameR1,text='计算结果：').pack(side='left')
        ttk.Button(self.FrameR1,text='刷新',width=4,command=self.fresh_data).pack(side='left')
        ttk.Button(self.FrameR1,text='保存结果',width=8,command=self.CommandButtonSaving).pack(side='left')
        self.DataD.set(np.round(0.5075,self.digit))
        self.DataDD.set(np.round(0.0035/0.5075*100,self.digit))
        FrameDataInput(self.FrameRes,'D',self.DataD,self.DataDD).pack(fill='x')
        self.Datarou.set(1000)
        FrameDataInput(self.FrameRes,'密度',self.Datarou,self.DataDrou).pack(fill='x')
        FrameDataTotal(self.FrameRes,'cd',self.Datacd,self.DataDcd).pack(fill='x')
        self.FrameR4=ttk.Frame(self.FrameRes)
        self.FrameR4.pack(fill='x')
        ttk.Label(self.FrameR4,text='(').pack(side='left')
        FrameDataTotal(self.FrameR4,'Pc',self.DataPccd,self.DataDPccd).pack(side='left')
        FrameDataTotal(self.FrameR4,'DP',self.DataDPcd,self.DataDDPcd).pack(side='left')
        ttk.Label(self.FrameR4,text=')').pack(side='left')
        FrameDataTotal(self.FrameRes,'U',self.DataU,self.DataDU).pack(fill='x')
        self.FrameR6=ttk.Frame(self.FrameRes)
        self.FrameR6.pack(fill='x')
        ttk.Label(self.FrameR6,text='(').pack(side='left')
        FrameDataTotal(self.FrameR6,'Pc',self.DataPc,self.DataDPc).pack(side='left')
        ttk.Label(self.FrameR6,text=')').pack(side='left')
    def fresh_data(self):
        try:
            self.data=np.genfromtxt(self.FDirData.get(), delimiter=',')
        except:
            self.data=np.genfromtxt('23_03_12_20_21_38_实验编号85.csv', delimiter=',')
        for i in self.Process:
            i.get_data(self.data)
            i.fresh_data()
            #TODO:刷新三幅图像的时间
        #计算cd和U
        self.DataPccd.set(self.Process[2].DataPc.get())
        self.DataDPccd.set(self.Process[2].DataDErrPPc.get())
        self.DataDPcd.set(self.Process[2].DataDP.get())
        self.DataDDPcd.set(self.Process[2].DataDErrPDP.get())
        self.Datacd.set(np.round(self.Process[2].DataQ.get()/(np.pi/4*self.DataD.get()**2*np.sqrt(2*self.Datarou.get()*self.DataDPcd.get()))/self.NumPipe.get(),self.digit))
        # self.DataDcd.set(np.round(np.sqrt(self.Process[2].DataDErrPQ.get()**2+4*self.DataDD.get()**2+1/4*(self.Process[2].DataDErrPDP.get()-self.Process[2].DataDDP.get())**2),self.digit))
        self.DataDcd.set(np.round(np.sqrt(self.Process[2].DataDErrPQ.get()**2+4*self.DataDD.get()**2+1/4*self.Process[2].DataDErrPDP.get()**2),self.digit))
        self.DataPc.set(self.Process[1].DataPc.get())
        self.DataDPc.set(self.Process[1].DataDErrPPc.get())
        self.DataU.set(np.round(self.Process[2].DataQ.get()/(np.pi/4*self.DataD.get()**2*np.sqrt(self.Process[1].DataDP.get()/self.Process[2].DataDP.get()))/self.NumPipe.get(),self.digit))
        self.DataDU.set(np.round(np.sqrt(self.Process[2].DataDErrPQ.get()**2+4*self.DataDD.get()**2+1/4*self.Process[1].DataDErrPDP.get()**2+1/4*self.Process[2].DataDErrPDP.get()**2),self.digit))
    def CommandButtonExpNoInc(self):
        self.ExpNo.set(self.ExpNo.get()+1)
        self.CommandButtonReadFDir()
    def CommandButtonExpNoDec(self):
        self.ExpNo.set(self.ExpNo.get()-1)
        self.CommandButtonReadFDir()
    def CommandButtonReadFDir(self):
        self.wfile=''
        for filename in os.listdir(self.DirData.get()):
            if '实验编号'+str(self.ExpNo.get())+'.csv' in filename:
                self.wfile=filename
        self.FDirData.set(self.DirData.get()+'/'+self.wfile)
        if self.wfile!='':
            self.fresh_data()
    def CommandButtonSaving(self):
        #TODO:ddd
        #data=['初始室压',p_0,'罐初始压力',pT_0,'路初始压力',pP_0,'室压',p,'罐压力',pT,'路压力',pP,'有效路压力',pPeff,'流量',Q,'实验段有效路压力',pexpeff]
        wb_target=openpyxl.load_workbook(self.FDirRes.get())
        ws_target=wb_target.active
        for i in range(1,ws_target.max_column+1):
            for k in ['U','Pc','cd','Pccd','DPcd']:
                if ws_target.cell(row=2,column=i).value==k:
                    ws_target.cell(self.ExpNo.get()+3,i,getattr(self,'Data'+k).get())
                    ws_target.cell(self.ExpNo.get()+3,i+1,getattr(self,'DataD'+k).get())
            if 'ExpIntStart' == ws_target.cell(row=2,column=i).value:
                ws_target.cell(self.ExpNo.get()+3,i,self.Process[1].IntL.get())
                ws_target.cell(self.ExpNo.get()+3,i+1,self.Process[1].IntR.get())
            if 'CdIntStart' == ws_target.cell(row=2,column=i).value:
                ws_target.cell(self.ExpNo.get()+3,i,self.Process[2].IntL.get())
                ws_target.cell(self.ExpNo.get()+3,i+1,self.Process[2].IntR.get())
        wb_target.save(self.FDirRes.get())
    def CommandButtonWriteRes(self):
        pass#TODO:写入
    def Update(self,name,index,mode):#更新数据库
        try:
            db.update(name,getattr(self,name).get())
        except:
            pass
    def new_var(self,name,type='String'):#后端，创建变量并与数据库同步
        setattr(self,name,getattr(tk,type+'Var')(name=name,value=db.get(name)))#后端
        getattr(self,name).trace_add('write',self.Update)#变量变动就马上更新数据库
def close_window():
    plt.close()
    root.quit()
    root.destroy()
if __name__=='__main__':
    root=tk.Tk()
    App=QProcess(root)
    App.pack()
    root.protocol("WM_DELETE_WINDOW",close_window)
    root.mainloop()