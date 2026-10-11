from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.chart import LineChart, Reference

from AtmProcess.core.excel_io import write_matrix
from AtmProcess.core.excel_workbook_coordinator import ExcelWorkbookCoordinator
from AtmProcess.domain.settings import AtomizationSettings
from AtmProcess.lib.dft import dft
from AtmProcess.services.atomization_angle_service import AtomizationBatchResult
from AtmProcess.services.summary_workbook_service import SummaryWorkbookService


@dataclass(frozen=True)
class AtomizationExportConfig:
    data_root: Path
    summary_path: Path
    max_freq: int
    experiment_name: str
    output_dir: Path | None = None
    experiment_no: int | None = None

    @property
    def experiment_dir(self) -> Path:
        return self.data_root / self.experiment_name

    @property
    def result_workbook(self) -> Path:
        if self.output_dir is None:
            return self.experiment_dir / "雾化锥角批处理结果.xlsx"
        return self.output_dir / f"{self.experiment_name}_雾化锥角批处理结果.xlsx"


class AtomizationResultExportService:
    def __init__(
        self,
        workbook_coordinator: ExcelWorkbookCoordinator | None = None,
        summary_service: SummaryWorkbookService | None = None,
    ) -> None:
        self.workbook_coordinator = workbook_coordinator or ExcelWorkbookCoordinator()
        self.summary_service = summary_service or SummaryWorkbookService(self.workbook_coordinator)

    def export(self, config: AtomizationExportConfig, settings: AtomizationSettings, result: AtomizationBatchResult) -> Path:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        self._write_batch_data(sheet, settings, result)
        angle_data = result.rows[:, 1:4]
        angle_mid = (angle_data[:, 0] + angle_data[:, 1]) / 2
        self._write_angle_mid(sheet, angle_mid)
        self._write_frequency_data(sheet, np.column_stack((angle_data, angle_mid)), config.max_freq, settings.frame_rate_hz)
        self._add_charts(sheet, config.max_freq)

        config.result_workbook.parent.mkdir(parents=True, exist_ok=True)
        result_workbook = self.workbook_coordinator.save_and_open(workbook, config.result_workbook)

        if config.summary_path.exists():
            self._update_summary(config, sheet)

        return result_workbook

    def _write_batch_data(self, sheet, settings: AtomizationSettings, result: AtomizationBatchResult) -> None:
        headers = [
            "XLeft", "XRight", "Y", "Rmax", "ThetaLeft", "ThetaRight", "window_length", "polyorder",
            "NumAvgL", "NumAvgR", "ThresRatioL", "ThresRatioR", "num_processes", "compute_time",
            "RStartL", "RStartR",
        ]
        values = [
            settings.x_left, settings.x_right, settings.y, settings.r_max, settings.theta_left, settings.theta_right,
            settings.window_length, settings.polyorder, settings.num_avg_left, settings.num_avg_right,
            settings.thres_ratio_left, settings.thres_ratio_right, result.process_count, result.compute_time,
            settings.r_start_left, settings.r_start_right,
        ]
        angle_data = result.rows[:, 1:4]
        means = np.mean(angle_data, axis=0)
        standard_deviations = np.std(angle_data, axis=0, ddof=1) if len(angle_data) > 1 else np.zeros(3)
        relative_deviations = standard_deviations / means[2]

        sheet.append(headers)
        sheet.append(values)
        sheet.append(["Frame", "AngleL", "AngleR", "AtmAngle"])
        sheet.append(["Average", *map(float, means)])
        sheet.append(["StDev", *map(float, standard_deviations)])
        sheet.append(["RelativeStDev", *map(float, relative_deviations)])
        for frame, angle_left, angle_right, atomization_angle in result.rows:
            sheet.append([int(frame), float(angle_left), float(angle_right), float(atomization_angle)])

    def _write_angle_mid(self, sheet, angle_mid: np.ndarray) -> None:
        standard_deviation = np.std(angle_mid, ddof=1) if len(angle_mid) > 1 else 0.0
        write_matrix(sheet, [["AngleM"]], row_start=3, col_start=5)
        write_matrix(sheet, [[np.mean(angle_mid)]], row_start=4, col_start=5)
        write_matrix(sheet, [[standard_deviation]], row_start=5, col_start=5)
        write_matrix(sheet, [[standard_deviation / sheet.cell(row=4, column=4).value]], row_start=6, col_start=5)
        for row_index, value in enumerate(angle_mid):
            sheet.cell(row=row_index + 7, column=5, value=value)

    def _write_frequency_data(self, sheet, result_matrix: np.ndarray, max_freq: int, frame_rate_hz: float) -> None:
        dft_result = (
            dft(result_matrix[:, 0], frame_rate_hz),
            dft(result_matrix[:, 1], frame_rate_hz),
            dft(result_matrix[:, 2], frame_rate_hz),
            dft(result_matrix[:, 3], frame_rate_hz),
        )
        freq_matrix = np.column_stack((
            dft_result[0][0],
            dft_result[0][1],
            dft_result[1][1],
            dft_result[2][1],
            dft_result[3][1],
        ))[:len(dft_result[0][0]) // 2, :]

        write_matrix(sheet, [["f/Hz", "左边界角幅值/°", "右边界角幅值/°", "雾化锥角幅值/°", "锥角中心线幅值/°"]], row_start=3, col_start=6)
        max_amplitudes = np.max(freq_matrix[1:, 1:], axis=0) if len(freq_matrix) > 1 else np.zeros(4)
        write_matrix(sheet, [["最大值", *max_amplitudes]], row_start=4, col_start=6)
        write_matrix(sheet, freq_matrix[:max_freq], row_start=5, col_start=6)

    def _add_charts(self, sheet, max_freq: int) -> None:
        legend_names = ["左锥角", "右锥角", "雾化锥角"]

        chart1 = LineChart()
        chart1.title = "时域信号"
        chart1.y_axis.title = "角度/°"
        chart1.x_axis.title = "帧数"
        chart1.add_data(Reference(sheet, min_col=2, min_row=7, max_col=5, max_row=sheet.max_row), titles_from_data=True)
        for index, name in enumerate(legend_names):
            chart1.series[index].name = name
        chart1.set_categories(Reference(sheet, min_col=1, min_row=7, max_row=sheet.max_row))
        chart1.width = 28
        chart1.height = 10
        sheet.add_chart(chart1, "K3")

        max_freq_row = max(5, min(max_freq, sheet.max_row))
        chart2 = LineChart()
        chart2.title = "频域信号：DFT分解"
        chart2.y_axis.title = "幅值/°"
        chart2.x_axis.title = "频率"
        chart2.add_data(Reference(sheet, min_col=7, min_row=5, max_col=10, max_row=max_freq_row), titles_from_data=True)
        for index, name in enumerate(legend_names):
            chart2.series[index].name = name
        chart2.set_categories(Reference(sheet, min_col=6, min_row=5, max_row=max_freq_row))
        chart2.width = 28
        chart2.height = 10
        sheet.add_chart(chart2, "K22")

        chart3 = LineChart()
        chart3.title = "频域信号：DFT分解"
        chart3.y_axis.title = "幅值/°"
        chart3.x_axis.title = "频率"
        chart3.add_data(Reference(sheet, min_col=7, min_row=5, max_col=10, max_row=max(5, sheet.max_row // 2)), titles_from_data=True)
        for index, name in enumerate(legend_names):
            chart3.series[index].name = name
        chart3.set_categories(Reference(sheet, min_col=6, min_row=5, max_row=max(5, sheet.max_row // 2)))
        chart3.width = 28
        chart3.height = 10
        sheet.add_chart(chart3, "K42")

    def _update_summary(self, config: AtomizationExportConfig, source_sheet) -> None:
        blocks = {
            "XLeft": [source_sheet.cell(2, column).value for column in range(1, 15)],
        }  # 旧参数块保持原顺序，避免把雾化参数 Y 写到液膜测量区的同名列
        for column in range(2, 6):
            title = source_sheet.cell(3, column).value
            blocks[title] = [source_sheet.cell(row, column).value for row in (4, 5)]
        extra_values = {
            source_sheet.cell(1, column).value: source_sheet.cell(2, column).value
            for column in (15, 16)
        }
        self.summary_service.update_blocks(
            config.summary_path,
            config.experiment_name,
            blocks,
            config.experiment_no,
            extra_values=extra_values,
        )
