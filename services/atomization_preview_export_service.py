from __future__ import annotations

from pathlib import Path

import openpyxl

from ImageProcessv3.core.excel_workbook_coordinator import ExcelWorkbookCoordinator


class AtomizationPreviewExportService:
    def __init__(self, workbook_coordinator: ExcelWorkbookCoordinator | None = None) -> None:
        self.workbook_coordinator = workbook_coordinator or ExcelWorkbookCoordinator()

    def export_gray_profile(self, folder: str | Path, res_left, res_left_sg, angle_left: float) -> Path:
        folder_path = Path(folder)
        save_path = folder_path / f"{folder_path.name}雾化角灰度分布.xlsx"

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.cell(row=1, column=1, value="resL")
        sheet.cell(row=1, column=3, value="resLSG")
        sheet.cell(row=1, column=4, value=angle_left)

        for index, data in enumerate([res_left, res_left_sg]):
            for row_index in range(len(data[0])):
                sheet.cell(row=row_index + 2, column=index * 2 + 1, value=data[0][row_index])
                sheet.cell(row=row_index + 2, column=index * 2 + 2, value=data[1][row_index])

        return self.workbook_coordinator.save_and_open(workbook, save_path)

    def export_profile(self, output_dir: str | Path, experiment_name: str, frame_no: int, profile_name: str, x_values, y_values, angle: float | None = None) -> Path:
        save_dir = Path(output_dir)
        save_dir.mkdir(parents=True, exist_ok=True)
        save_path = save_dir / f"{experiment_name}_{frame_no}_{profile_name}.xlsx"

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "灰度分布"
        sheet.append(["角度/°", "面积加权平均灰度值"])
        if angle is not None:
            sheet.cell(row=1, column=4, value="提取角度/°")
            sheet.cell(row=1, column=5, value=angle)
        for x_value, y_value in zip(x_values, y_values):
            sheet.append([float(x_value), float(y_value)])

        return self.workbook_coordinator.save_and_open(workbook, save_path)
