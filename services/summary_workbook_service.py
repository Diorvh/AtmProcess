from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

import openpyxl


class SummaryWorkbookService:
    experiment_headers = ("编号", "实验编号", "序号", "No")
    header_aliases = {
        "像素大小/um/pix": ("像素大小/um/pix", "像素尺寸/um/pix", "PixSize"),
        "平均液膜宽度/mm": ("平均液膜宽度/mm",),
        "平均液膜长度/mm": ("平均液膜长度/mm",),
        "平均破碎长度/mm": ("平均破碎长度/mm",),
        "半数采样破碎长度/mm": ("半数采样破碎长度/mm",),
        "总张数": ("总张数",),
        "起始帧": ("起始帧",),
        "结束帧": ("结束帧",),
    }

    def update(self, summary_path: str | Path, experiment_name: str, values: Mapping[str, object], experiment_no: int | None = None) -> int:
        workbook = openpyxl.load_workbook(Path(summary_path))
        sheet = workbook.active
        headers = self._header_map(sheet)
        target_row = self._find_experiment_row(sheet, headers, experiment_name, experiment_no)

        written_count = 0
        for title, value in values.items():
            column = self._find_column(headers, title)
            if column is None:
                continue
            sheet.cell(target_row, column, value)  # 按第一行标题写入，兼容不同汇总表的列顺序
            written_count += 1

        workbook.save(summary_path)
        return written_count

    def update_contiguous_block(
        self,
        summary_path: str | Path,
        experiment_name: str,
        anchor_title: str,
        anchor_offset: int,
        values: Sequence[object],
        experiment_no: int | None = None,
    ) -> int:
        workbook = openpyxl.load_workbook(Path(summary_path))
        sheet = workbook.active
        headers = self._header_map(sheet)
        target_row = self._find_experiment_row(sheet, headers, experiment_name, experiment_no)
        anchor_column = self._find_column(headers, anchor_title)
        if anchor_column is None:
            return 0

        start_column = anchor_column - anchor_offset
        if start_column < 1:
            return 0

        for index, value in enumerate(values):
            sheet.cell(target_row, start_column + index, value)  # 连续参数块只定位起始列，然后按顺序写入

        workbook.save(summary_path)
        return len(values)

    def read_contiguous_block(
        self,
        summary_path: str | Path,
        experiment_name: str,
        anchor_title: str,
        count: int,
        experiment_no: int | None = None,
    ) -> list[object]:
        workbook = openpyxl.load_workbook(Path(summary_path), data_only=True)
        sheet = workbook.active
        headers = self._header_map(sheet)
        target_row = self._find_existing_experiment_row(sheet, headers, experiment_name, experiment_no)
        if target_row is None:
            return []
        anchor_column = self._find_column(headers, anchor_title)
        if anchor_column is None:
            return []
        return [sheet.cell(target_row, anchor_column + index).value for index in range(count)]

    def _header_map(self, sheet) -> dict[str, int]:
        headers: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(1, column).value
            if value is not None and str(value).strip():
                headers.setdefault(str(value).strip(), column)  # 同名标题保留第一次出现的位置，例如破碎长度区和雾化区都有 Y
        return headers

    def _find_column(self, headers: Mapping[str, int], title: str) -> int | None:
        for alias in (title, *self.header_aliases.get(title, ())):
            column = headers.get(alias)
            if column is not None:
                return column
        return None

    def _find_experiment_row(self, sheet, headers: Mapping[str, int], experiment_name: str, experiment_no: int | None) -> int:
        id_column = self._find_identifier_column(headers)
        if id_column is not None:
            for row in range(2, sheet.max_row + 1):
                if self._matches_experiment(sheet.cell(row, id_column).value, experiment_name, experiment_no):
                    return row
            row = sheet.max_row + 1
            sheet.cell(row, id_column, experiment_no if experiment_no is not None else experiment_name)  # 找不到时追加新实验行
            return row

        for row in range(2, sheet.max_row + 1):
            for column in range(1, sheet.max_column + 1):
                if self._matches_experiment(sheet.cell(row, column).value, experiment_name, experiment_no):
                    return row
        return sheet.max_row + 1

    def _find_existing_experiment_row(self, sheet, headers: Mapping[str, int], experiment_name: str, experiment_no: int | None) -> int | None:
        id_column = self._find_identifier_column(headers)
        if id_column is not None:
            for row in range(2, sheet.max_row + 1):
                if self._matches_experiment(sheet.cell(row, id_column).value, experiment_name, experiment_no):
                    return row
            return None

        for row in range(2, sheet.max_row + 1):
            for column in range(1, sheet.max_column + 1):
                if self._matches_experiment(sheet.cell(row, column).value, experiment_name, experiment_no):
                    return row
        return None

    def _find_identifier_column(self, headers: Mapping[str, int]) -> int | None:
        for title in self.experiment_headers:
            if title in headers:
                return headers[title]
        return None

    def _matches_experiment(self, value, experiment_name: str, experiment_no: int | None) -> bool:
        if value is None:
            return False
        if experiment_no is not None and isinstance(value, (int, float)):
            return int(value) == experiment_no
        text = str(value).strip()
        candidates = {experiment_name}
        if experiment_no is not None:
            candidates.update({str(experiment_no), f"实验编号{experiment_no}"})
        return text in candidates
