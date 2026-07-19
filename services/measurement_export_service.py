from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import numpy as np
import openpyxl

from ImageProcessv3.core.excel_io import write_matrix
from ImageProcessv3.core.excel_workbook_coordinator import ExcelWorkbookCoordinator
from ImageProcessv3.services.summary_workbook_service import SummaryWorkbookService


@dataclass(frozen=True)
class MeasurementRecordImport:
    mode: str
    rows: list[list]


class MeasurementExportService:
    sheet_headers = ["帧数", "左线", "右线", "上线", "下线", "宽度/pix", "宽度/mm", "长度/pix", "长度/mm"]
    breakup_headers = ["帧数", "破碎线", "破碎长度/pix", "破碎长度/mm"]

    def __init__(
        self,
        summary_service: SummaryWorkbookService | None = None,
        workbook_coordinator: ExcelWorkbookCoordinator | None = None,
    ) -> None:
        self.workbook_coordinator = workbook_coordinator or ExcelWorkbookCoordinator()
        self.summary_service = summary_service or SummaryWorkbookService(self.workbook_coordinator)

    def read_records(self, file_path: str | Path) -> MeasurementRecordImport:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        try:
            sheet = workbook.active
            breakup_headers = [sheet.cell(2, column).value for column in range(1, len(self.breakup_headers) + 1)]
            mode = "breakup" if breakup_headers == self.breakup_headers else "sheet"
            column_count = len(self.breakup_headers) if mode == "breakup" else len(self.sheet_headers)
            rows = []
            for values in sheet.iter_rows(min_row=3, max_col=column_count, values_only=True):
                if values[0] is None:
                    continue
                rows.append(self._breakup_record(values) if mode == "breakup" else self._sheet_record(values))
        finally:
            workbook.close()
        return MeasurementRecordImport(mode, rows)

    def export_sheet(
        self,
        folder: str | Path,
        experiment_name: str,
        rows: Sequence[Sequence],
        impact_x: int,
        impact_y: int,
        pixel_size_um: float,
        output_dir: str | Path | None = None,
        summary_path: str | Path | None = None,
        experiment_no: int | None = None,
        decimate: int = 1,
    ) -> Path:
        save_path = self._save_path(folder, "液膜信息", experiment_name, "液膜边界测量.xlsx", output_dir)
        numeric = np.array([row[2:] for row in rows], dtype=float)
        frames = np.array([row[1] for row in rows], dtype=float)
        width_mm = np.array(numeric[:, 5], dtype=float)
        length_mm = np.array(numeric[:, 7], dtype=float)
        average_width = float(np.mean(width_mm))
        average_length = float(np.mean(length_mm))
        title = [
            "X",
            impact_x,
            "Y",
            impact_y,
            "像素大小/um/pix",
            pixel_size_um,
            "Decimate",
            decimate,
            "总张数",
            len(rows),
            "起始帧",
            int(frames[0]),
            "结束帧",
            int(frames[-1]),
        ]
        stats = [
            ("平均液膜宽度/mm", average_width),
            ("液膜宽度标准差/mm", self._sample_std(width_mm)),
            ("平均液膜长度/mm", average_length),
            ("液膜长度标准差/mm", self._sample_std(length_mm)),
        ]
        self._write(save_path, title, self.sheet_headers, self._without_image_name(rows), stats)
        self._update_summary(
            summary_path,
            experiment_name,
            {
                "X": impact_x,
                "Y": impact_y,
                "像素大小/um/pix": pixel_size_um,
                "平均液膜宽度/mm": average_width,
                "平均液膜长度/mm": average_length,
            },
            experiment_no,
        )
        return save_path

    def export_breakup(
        self,
        folder: str | Path,
        experiment_name: str,
        rows: Sequence[Sequence],
        impact_x: int,
        impact_y: int,
        pixel_size_um: float,
        decimate: int,
        output_dir: str | Path | None = None,
        summary_path: str | Path | None = None,
        experiment_no: int | None = None,
    ) -> Path:
        save_path = self._save_path(folder, "破碎长度", experiment_name, "破碎长度测量.xlsx", output_dir)
        numeric = np.array(rows, dtype=object)
        frames = np.array(numeric[:, 1], dtype=float)
        lengths = np.array(numeric[:, 4], dtype=float)
        average_breakup_length = self._frame_weighted_average(frames, lengths)
        title = [
            "X",
            impact_x,
            "Y",
            impact_y,
            "像素大小/um/pix",
            pixel_size_um,
            "Decimate",
            decimate,
            "总张数",
            len(rows),
            "起始帧",
            int(frames[0]),
            "结束帧",
            int(frames[-1]),
        ]
        stats = [
            ("平均破碎长度/mm", average_breakup_length),
            ("破碎长度标准差/mm", self._sample_std(lengths)),
        ]
        self._write(save_path, title, self.breakup_headers, self._without_image_name(rows), stats)
        self._update_breakup_summary(
            summary_path,
            experiment_name,
            [impact_x, impact_y, pixel_size_um, average_breakup_length, decimate, len(rows), int(frames[0]), int(frames[-1])],
            experiment_no,
        )
        return save_path

    def _save_path(self, folder: str | Path, subfolder: str, experiment_name: str, suffix: str, output_dir: str | Path | None) -> Path:
        folder_path = Path(folder)
        save_dir = Path(output_dir) if output_dir else folder_path.parent / subfolder
        save_dir.mkdir(parents=True, exist_ok=True)
        return save_dir / f"{experiment_name}_{suffix}"

    def _write(self, save_path: Path, title: list, headers: list[str], rows: Sequence[Sequence], stats: Sequence[tuple[str, float]]) -> Path:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        write_matrix(sheet, [title])
        stats_headers = [name for name, _value in stats]
        write_matrix(sheet, [headers + [""] + stats_headers], append=True)
        stats_values = [value for _name, value in stats]
        data_rows = []
        for index, row in enumerate(rows):
            row_stats = stats_values if index == 0 else [""] * len(stats_values)
            data_rows.append(list(row) + [""] + row_stats)
        write_matrix(sheet, data_rows, append=True)
        return self.workbook_coordinator.save_and_open(workbook, save_path)

    def _update_summary(self, summary_path: str | Path | None, experiment_name: str, values: dict[str, object], experiment_no: int | None) -> None:
        if not summary_path:
            return
        self.summary_service.update(summary_path, experiment_name, values, experiment_no)

    def _update_breakup_summary(self, summary_path: str | Path | None, experiment_name: str, values: list, experiment_no: int | None) -> None:
        if not summary_path:
            return
        self.summary_service.update_contiguous_block(summary_path, experiment_name, "平均破碎长度/mm", 3, values, experiment_no)

    def _frame_weighted_average(self, frames: np.ndarray, lengths: np.ndarray) -> float:
        if len(frames) <= 1:
            return float(lengths[0])
        area = 0.0
        for index in range(len(frames) - 1):
            area += (frames[index + 1] - frames[index] + 1) * (lengths[index + 1] + lengths[index]) / 2
        area = area - np.sum(lengths) + lengths[0] + lengths[-1]
        return float(area / (frames[-1] - frames[0] + 1))

    def _sample_std(self, values: np.ndarray) -> float:
        if len(values) <= 1:
            return 0.0
        return float(np.std(values, ddof=1))

    def _without_image_name(self, rows: Sequence[Sequence]) -> list[list]:
        return [list(row[1:]) for row in rows]

    @staticmethod
    def _sheet_record(values: Sequence[object]) -> list:
        frame = int(values[0])
        return [
            str(frame),
            frame,
            int(values[1]),
            int(values[2]),
            int(values[3]),
            int(values[4]),
            int(values[5]),
            float(values[6]),
            int(values[7]),
            float(values[8]),
        ]

    @staticmethod
    def _breakup_record(values: Sequence[object]) -> list:
        frame = int(values[0])
        return [str(frame), frame, int(values[1]), int(values[2]), float(values[3])]
